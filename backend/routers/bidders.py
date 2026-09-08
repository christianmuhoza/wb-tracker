"""
Bidder router — extracts all bidder endpoints from the original monolithic api.py.
"""

import csv
import io
import json
import re
import threading
import unicodedata
from datetime import datetime
from typing import Any

import openpyxl
from auth import require_auth
from db import db, q
from fastapi import APIRouter, Depends, HTTPException, Query
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from ratelimit import STRICT_LIMITER, rate_limited
from services.bidder_extraction import (
    _infer_notice_category,
    _merge_categories,
    extract_awarded_bidders,
    extract_bidders_list,
    fetch_bidders_from_detail_page,
    parse_notice_bidder_details,
)
from services.contact_enrichment import search_company_contact
from services.gemini_service import classify_and_enrich_with_gemini
from services.tech import build_tech_bidder_condition, build_tech_notice_condition, looks_like_tech_bidder

router = APIRouter(prefix="/api/bidders", tags=["bidders"])


# ── Pydantic models ──────────────────────────────────────────────────────────


class BidderUpdate(BaseModel):
    name: str | None = None
    category: str | None = None
    contact_name: str | None = None
    contact_email: str | None = None
    contact_phone: str | None = None
    linkedin_url: str | None = None
    contact_org: str | None = None
    country: str | None = None
    business_model: str | None = None
    core_products: str | None = None
    corporate_activities: str | None = None


class EnrichResult(BaseModel):
    status: str
    bidder_id: int
    bidder_name: str
    found: dict[str, Any] = {}
    updated_fields: list[str] = []
    error: str | None = None


# ── Country-normalisation helpers ─────────────────────────────────────────────

BIDDER_COUNTRY_ALIASES = {
    "Benin": ["Bénin", "B?nin"],
    "Central African Republic": ["République centrafricaine", "R?publique centrafricaine"],
    "Gambia": ["Gambia, The"],
    "Guinea": ["Guinée"],
    "Mozambique": ["Moçambique", "Mo?ambique"],
    "Tanzania": ["Tanzánia"],
    "Somalia, Federal Republic of": ["Somalia"],
}

_ACCENT_FROM = "ÁÀÂÃÄÅÇÈÉÊËÍÌÎÏÑÓÒÔÕÖÙÚÛÜÝáàâãäåçèéêëíìîïñóòôõöùúûüýÿ"
_ACCENT_TO = "AAAAAACEEEEIIIINOOOOOUUUUYaaaaaaceeeeiiiinooooouuuuyy"
_CORRUPT_CHARS = "\ufffd?\"'" + ","


def _country_match_keys(country):
    """Accent/case/mojibake-insensitive keys for a country name and its aliases."""
    names = {country, *BIDDER_COUNTRY_ALIASES.get(country, [])}
    keys = set()
    for name in names:
        name = unicodedata.normalize("NFKD", name)
        name = "".join(ch for ch in name if not unicodedata.combining(ch))
        keys.add("".join(ch for ch in name.lower() if ch not in _CORRUPT_CHARS))
    return sorted(keys)


def _bidder_country_clause(country, prefix="b"):
    col = f"{prefix}.country"
    exact = [country, *BIDDER_COUNTRY_ALIASES.get(country, [])]
    norm_expr = f"lower(translate({col}, %s, %s))"
    return (
        f"({col} = ANY(%s::text[]) OR {norm_expr} = ANY(%s::text[]))",
        [exact, _ACCENT_FROM, _ACCENT_TO, _country_match_keys(country)],
    )


def build_bidder_filters(search: str | None, qs: str | None, country: str | None, tech_only: bool = False):
    search_term = search or qs
    filters = ["1=1"]
    params: list[Any] = []
    if search_term:
        filters.append("(b.name ILIKE %s OR b.contact_org ILIKE %s OR b.contact_name ILIKE %s)")
        like = f"%{search_term}%"
        params.extend([like, like, like])
    if country:
        clause, country_params = _bidder_country_clause(country)
        filters.append(clause)
        params.extend(country_params)
    if tech_only:
        bidder_filter, bidder_params = build_tech_bidder_condition("b")
        notice_filter, notice_params = build_tech_notice_condition("pn")
        filters.append(f"""(
            {bidder_filter}
            OR EXISTS (
                SELECT 1
                FROM bidder_awards tech_ba
                JOIN procurement_notices pn ON pn.id = tech_ba.notice_id
                WHERE tech_ba.bidder_id = b.id
                  AND {notice_filter}
            )
        )""")
        params.extend(bidder_params + notice_params)
    return " AND ".join(filters), params


# ── Export helpers ────────────────────────────────────────────────────────────

BIDDER_EXPORT_FIELDS = {
    "c": ("#", "c", 5),
    "name": ("Company", "name", 40),
    "country": ("Supplier Country", "country", 20),
    "project_country": ("Project Country / Region", "project_country", 22),
    "project_id": ("Project ID", "project_id", 18),
    "contract_reference": ("Contract Reference", "contract_reference", 20),
    "title": ("Contract / Scope", "title", 50),
    "award_date": ("Award Date", "award_date", 15),
    "currency": ("Currency", "currency", 10),
    "contract_value": ("Contract Value", "contract_value", 20),
    "company_website": ("Company Website", "company_website", 25),
    "url": ("World Bank Award Notice", "url", 25),
    "notes": ("Notes", "notes", 30),
}


def resolve_bidder_export_fields(fields: str | None) -> list[str]:
    selected_fields = [f.strip() for f in (fields or "").split(",") if f.strip()]
    if not selected_fields:
        return [
            "c",
            "name",
            "country",
            "project_country",
            "project_id",
            "contract_reference",
            "title",
            "award_date",
            "currency",
            "contract_value",
            "company_website",
            "url",
            "notes",
        ]
    return [field for field in selected_fields if field in BIDDER_EXPORT_FIELDS]


def fetch_bidder_export_rows(
    search: str | None, qs: str | None, country: str | None, won_only: bool, tech_only: bool = False
):
    where, params = build_bidder_filters(search, qs, country, tech_only)
    rows = q(
        f"""
        SELECT
            b.name,
            b.country,
            w.project_country,
            w.project_id,
            w.contract_reference,
            w.title,
            w.award_date,
            w.currency,
            w.contract_value,
            w.url
        FROM bidders b
        LEFT JOIN LATERAL (
            SELECT
                pn.country   AS project_country,
                pn.project_id,
                pn.borrower_bid_reference AS contract_reference,
                pn.title,
                ba.award_date::text AS award_date,
                ba.currency,
                ba.award_amount AS contract_value,
                CASE
                    WHEN pn.url IS NOT NULL AND pn.url <> '' AND pn.url <> 'https://projects.worldbank.org/en/projects-operations/procurement'
                    THEN pn.url
                    ELSE 'https://projects.worldbank.org/en/projects-operations/procurement-detail/' || pn.id
                END AS url
            FROM bidder_awards ba
            JOIN procurement_notices pn ON pn.id = ba.notice_id
            WHERE ba.bidder_id = b.id AND ba.won = TRUE
            ORDER BY ba.award_date DESC NULLS LAST
            LIMIT 1
        ) w ON TRUE
        WHERE {where}
        ORDER BY b.name ASC
    """,
        params,
    )
    if won_only:
        rows = [r for r in rows if r.get("award_date")]
    for idx, row in enumerate(rows, 1):
        row["c"] = idx
        row["company_website"] = None
        row["notes"] = None
    return rows


def _write_bidder_export_sheet(ws, selected_fields: list[str], rows) -> None:
    for col_idx, field in enumerate(selected_fields, 1):
        label, _, width = BIDDER_EXPORT_FIELDS[field]
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5F43")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, row in enumerate(rows, 2):
        for col_idx, field in enumerate(selected_fields, 1):
            value = row.get(BIDDER_EXPORT_FIELDS[field][1])
            ws.cell(row=row_idx, column=col_idx, value=value)


def _safe_sheet_title(title: str, fallback: str) -> str:
    safe = re.sub(r"[\\/*?\[\]:]", "_", title).strip()[:31]
    return safe or fallback


# ── Core bidder extraction pipeline ───────────────────────────────────────────


@router.post("/from_notice")
def import_bidders_from_notice(
    notice_id: str = Query(None), fetch_detail: bool = Query(False), _: dict = Depends(require_auth)
):
    """
    Extract bidders from a stored notice and upsert into `bidders` and `bidder_awards`.

    IMPORTANT: The `borrower` field is the HOST INSTITUTION (government client, e.g.
    "Ministry of Finance"). It is NEVER a bidder. Only companies found in the notice
    description text or scraped from the WB detail page are stored as bidders.
    """
    if not notice_id:
        raise HTTPException(status_code=400, detail="notice_id is required")

    rows = q("SELECT * FROM procurement_notices WHERE id = %s", [notice_id])
    if not rows:
        raise HTTPException(status_code=404, detail="Notice not found")

    notice = rows[0]
    desc = notice.get("description") or ""
    bidder_details = parse_notice_bidder_details(desc)
    details_by_name = {detail["name"].lower(): detail for detail in bidder_details if detail.get("name")}
    inferred_category = _infer_notice_category(notice, desc)

    # Extract all participating bidders and explicitly awarded winners
    all_bidders = [detail["name"] for detail in bidder_details] or extract_bidders_list(desc)
    awarded = [
        detail["name"] for detail in bidder_details if detail.get("section") == "awarded"
    ] or extract_awarded_bidders(desc)

    # Optionally scrape the WB detail page for richer bidder data
    if fetch_detail:
        detail_bidders = fetch_bidders_from_detail_page(notice_id, notice.get("url") or "")
        for b in detail_bidders:
            if b.lower() not in {x.lower() for x in all_bidders}:
                all_bidders.append(b)

    # If nothing found — return cleanly. Do NOT fall back to borrower.
    # The borrower is the CLIENT, not a bidding company.
    if not all_bidders:
        return {
            "status": "ok",
            "bidders_found": 0,
            "bidders_inserted": 0,
            "links_created": 0,
            "note": (
                "No bidder names found in the description or detail page. "
                "This notice may not contain structured bidder data. "
                "Try 'Fetch from World Bank Page' for richer data."
            ),
        }

    inserted = 0
    linked = 0

    with db() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT bidder_id FROM bidder_awards WHERE notice_id = %s", (notice_id,))
            previous_bidder_ids = [
                row["bidder_id"] if hasattr(row, "__getitem__") else row[0] for row in cur.fetchall()
            ]
            cur.execute("DELETE FROM bidder_awards WHERE notice_id = %s", (notice_id,))

            for name in all_bidders:
                detail = details_by_name.get(name.lower(), {})
                bidder_country = detail.get("country")
                bidder_category = inferred_category
                bid_amount = detail.get("amount")
                bid_currency = detail.get("currency") or notice.get("currency")

                # Upsert bidder
                cur.execute("SELECT id, category, country FROM bidders WHERE lower(name) = lower(%s)", (name,))
                existing = cur.fetchone()
                if existing:
                    bidder_id = existing["id"] if hasattr(existing, "__getitem__") else existing[0]
                    existing_category = existing.get("category") if hasattr(existing, "get") else None
                    cur.execute(
                        """UPDATE bidders
                           SET category = %s,
                               country = COALESCE(NULLIF(country, ''), %s),
                               updated_at = NOW()
                           WHERE id = %s""",
                        (_merge_categories(existing_category, bidder_category), bidder_country, bidder_id),
                    )
                else:
                    cur.execute(
                        """INSERT INTO bidders (name, category, country, created_at, updated_at)
                           VALUES (%s, %s, %s, NOW(), NOW())
                           RETURNING id""",
                        (name, bidder_category, bidder_country),
                    )
                    row = cur.fetchone()
                    bidder_id = row["id"] if hasattr(row, "__getitem__") else row[0]
                    inserted += 1

                # won=True only when explicitly named as awarded in the text
                # Never derived from the borrower field
                won = detail.get("section") == "awarded" or any(name.lower() == w.lower() for w in awarded)

                # Link to notice
                cur.execute(
                    "SELECT id FROM bidder_awards WHERE bidder_id = %s AND notice_id = %s", (bidder_id, notice_id)
                )
                link = cur.fetchone()
                if link:
                    link_id = link["id"] if hasattr(link, "__getitem__") else link[0]
                    cur.execute(
                        """UPDATE bidder_awards
                           SET won = %s, award_amount = %s, currency = %s,
                               award_date = %s, role = %s, updated_at = NOW()
                           WHERE id = %s""",
                        (
                            won,
                            bid_amount or notice.get("contract_amount"),
                            bid_currency,
                            notice.get("notice_date"),
                            detail.get("role"),
                            link_id,
                        ),
                    )
                else:
                    cur.execute(
                        """INSERT INTO bidder_awards
                           (bidder_id, notice_id, won, award_amount, currency,
                            award_date, role, created_at, updated_at)
                           VALUES (%s, %s, %s, %s, %s, %s, %s, NOW(), NOW())""",
                        (
                            bidder_id,
                            notice_id,
                            won,
                            bid_amount or notice.get("contract_amount"),
                            bid_currency,
                            notice.get("notice_date"),
                            detail.get("role"),
                        ),
                    )
                    linked += 1

            if previous_bidder_ids:
                cur.execute(
                    """DELETE FROM bidders b
                       WHERE b.id = ANY(%s)
                         AND NOT EXISTS (
                             SELECT 1 FROM bidder_awards ba
                             WHERE ba.bidder_id = b.id
                         )""",
                    (previous_bidder_ids,),
                )

        conn.commit()

    return {
        "status": "ok",
        "bidders_found": len(all_bidders),
        "bidders_inserted": inserted,
        "links_created": linked,
        "awarded_identified": len(awarded),
    }


# ── List bidders with aggregated stats ────────────────────────────────────────


@router.get("/")
def list_bidders(
    qs: str | None = Query(None, description="Search by name (legacy param)"),
    search: str | None = Query(None, description="Search by name or org"),
    country: str | None = Query(None, description="Filter by country"),
    won_only: bool = Query(False, description="Only bidders who won at least one award"),
    tech_only: bool = Query(False, description="Only tech-related bidders"),
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=200),
    sort_by: str = Query("bid_count", regex="^(bid_count|won_count|total_bid_amount|last_bid_date|name)$"),
    sort_order: str = Query("desc", regex="^(asc|desc)$"),
):
    """List bidders with aggregated bid stats."""
    where, params = build_bidder_filters(search, qs, country, tech_only)
    having = "HAVING COUNT(CASE WHEN ba.won THEN 1 END) > 0" if won_only else ""
    offset = (page - 1) * page_size
    sort_exprs = {
        "bid_count": "bid_count",
        "won_count": "won_count",
        "total_bid_amount": "total_bid_amount",
        "last_bid_date": "last_bid_date",
        "name": "LOWER(b.name)",
    }
    order_dir = "ASC" if sort_order == "asc" else "DESC"
    nulls = "NULLS FIRST" if order_dir == "ASC" else "NULLS LAST"
    order_expr = sort_exprs[sort_by]
    tie_breaker = "bid_count DESC, b.name ASC" if sort_by != "name" else "bid_count DESC"
    tech_notice_filter, tech_notice_params = build_tech_notice_condition("pn")

    rows = q(
        f"""
        SELECT
            b.id,
            b.name,
            b.category,
            b.contact_name,
            b.contact_email,
            b.contact_phone,
            b.contact_org,
            b.country,
            b.business_model,
            b.core_products,
            b.corporate_activities,
            b.created_at,
            b.updated_at,
            COUNT(ba.id)                              AS bid_count,
            COUNT(CASE WHEN ba.won THEN 1 END)        AS won_count,
            COALESCE(SUM(ba.award_amount), 0)         AS total_bid_amount,
            MAX(NULLIF(ba.currency, ''))              AS primary_currency,
            MAX(ba.award_date)::text                  AS last_bid_date,
            (SELECT COUNT(DISTINCT pn.id)
             FROM bidder_awards ba3
             JOIN procurement_notices pn ON pn.id = ba3.notice_id
             WHERE ba3.bidder_id = b.id
               AND {tech_notice_filter})              AS tech_notice_count,
            (SELECT pn.title
             FROM bidder_awards ba2
             JOIN procurement_notices pn ON pn.id = ba2.notice_id
             WHERE ba2.bidder_id = b.id
             ORDER BY ba2.award_date DESC NULLS LAST
             LIMIT 1)                                 AS latest_bid_title
        FROM bidders b
        LEFT JOIN bidder_awards ba ON ba.bidder_id = b.id
        WHERE {where}
        GROUP BY b.id
        {having}
        ORDER BY {order_expr} {order_dir} {nulls}, {tie_breaker}
        LIMIT %s OFFSET %s
    """,
        tech_notice_params + params + [page_size, offset],
    )

    for row in rows:
        row["is_tech"] = looks_like_tech_bidder(row) or int(row.get("tech_notice_count") or 0) > 0

    total_rows = q(
        f"""
        SELECT COUNT(*) AS cnt FROM (
            SELECT b.id
            FROM bidders b
            LEFT JOIN bidder_awards ba ON ba.bidder_id = b.id
            WHERE {where}
            GROUP BY b.id
            {having}
        ) sub
    """,
        params,
    )
    total = total_rows[0]["cnt"] if total_rows else 0

    return {
        "page": page,
        "page_size": page_size,
        "total": total,
        "total_pages": max(1, -(-total // page_size)),
        "data": [dict(r) for r in rows],
    }


# ── Country list ──────────────────────────────────────────────────────────────


@router.get("/countries")
def list_bidder_countries():
    rows = q("""
        SELECT DISTINCT country
        FROM bidders
        WHERE country IS NOT NULL AND TRIM(country) <> ''
        ORDER BY country
    """)
    return [row["country"] for row in rows]


# ── Import status ─────────────────────────────────────────────────────────────


@router.get("/import_status")
def get_bidder_import_status():
    total_awards = q(
        "SELECT COUNT(*) AS cnt FROM procurement_notices WHERE notice_type IN ('Contract Award','Award') OR notice_type ILIKE '%Contract Award%' OR status = 'Awarded'"
    )
    linked_awards = q("SELECT COUNT(DISTINCT notice_id) AS cnt FROM bidder_awards")
    missing_awards = q("""
        SELECT COUNT(*) AS cnt
        FROM procurement_notices pn
        WHERE (pn.notice_type IN ('Contract Award','Award') OR pn.notice_type ILIKE '%Contract Award%' OR pn.status = 'Awarded')
          AND NOT EXISTS (SELECT 1 FROM bidder_awards ba WHERE ba.notice_id = pn.id)
    """)
    sample_missing = q("""
        SELECT pn.id, pn.title, pn.notice_date::text
        FROM procurement_notices pn
        WHERE (pn.notice_type IN ('Contract Award','Award') OR pn.notice_type ILIKE '%Contract Award%' OR pn.status = 'Awarded')
          AND NOT EXISTS (SELECT 1 FROM bidder_awards ba WHERE ba.notice_id = pn.id)
        ORDER BY pn.notice_date DESC NULLS LAST
        LIMIT 10
    """)

    summary = None
    try:
        with open("bidders_import_summary.json", encoding="utf-8") as f:
            summary = json.loads(f.read())
    except Exception:
        summary = None

    return {
        "total_award_notices": total_awards[0]["cnt"] if total_awards else 0,
        "linked_award_notices": linked_awards[0]["cnt"] if linked_awards else 0,
        "missing_award_notices": missing_awards[0]["cnt"] if missing_awards else 0,
        "missing_samples": [dict(r) for r in sample_missing],
        "last_summary": summary,
    }


# ── Single bidder CRUD ────────────────────────────────────────────────────────


@router.get("/{bidder_id}")
def get_bidder(bidder_id: int):
    rows = q("SELECT * FROM bidders WHERE id = %s", [bidder_id])
    if not rows:
        raise HTTPException(status_code=404, detail="Bidder not found")
    return dict(rows[0])


@router.get("/{bidder_id}/notices")
def get_bidder_notices(bidder_id: int):
    """All procurement notices a bidder participated in, with bid amounts and win status."""
    rows = q(
        """
        SELECT
            pn.id,
            pn.title,
            pn.project_name,
            pn.project_id,
            pn.country          AS borrower_country,
            pn.notice_type,
            pn.notice_date::text,
            COALESCE(pn.submission_date, pn.submission_deadline::date)::text AS submission_date,
            pn.status,
            pn.url,
            pn.contract_amount  AS notice_contract_amount,
            pn.currency         AS notice_currency,
            pn.borrower,
            ba.won,
            ba.award_amount,
            ba.currency,
            COALESCE(ba.award_amount, pn.contract_amount) AS bid_amount,
            COALESCE(NULLIF(ba.currency, ''), NULLIF(pn.currency, '')) AS bid_currency,
            ba.award_date::text,
            ba.role
        FROM bidder_awards ba
        JOIN procurement_notices pn ON pn.id = ba.notice_id
        WHERE ba.bidder_id = %s
        ORDER BY ba.award_date DESC NULLS LAST, pn.notice_date DESC NULLS LAST
    """,
        [bidder_id],
    )
    return [dict(r) for r in rows]


@router.put("/{bidder_id}")
def update_bidder(bidder_id: int, body: BidderUpdate, _: dict = Depends(require_auth)):
    fields = {k: v for k, v in body.dict().items() if v is not None}
    if not fields:
        return {"status": "ok"}
    sets = []
    params = []
    for k, v in fields.items():
        sets.append(f"{k} = %s")
        params.append(v)
    params.append(bidder_id)
    sql = f"UPDATE bidders SET {', '.join(sets)}, updated_at = NOW() WHERE id = %s"
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(sql, params)
        conn.commit()
    return {"status": "ok"}


@router.delete("/{bidder_id}")
def delete_bidder(bidder_id: int, _: dict = Depends(require_auth)):
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id, name FROM bidders WHERE id = %s", [bidder_id])
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="Bidder not found")
            cur.execute("DELETE FROM bidders WHERE id = %s", [bidder_id])
        conn.commit()
    return {"status": "ok", "deleted_id": bidder_id, "deleted_name": row["name"]}


# ── Batch import ──────────────────────────────────────────────────────────────


@router.post("/import_missing")
def import_missing_awards(fetch_detail: bool = Query(False), _: dict = Depends(require_auth)):
    notices = q("""
        SELECT pn.id
        FROM procurement_notices pn
        WHERE (pn.notice_type IN ('Contract Award','Award') OR pn.notice_type ILIKE '%Contract Award%' OR pn.status = 'Awarded')
          AND NOT EXISTS (SELECT 1 FROM bidder_awards ba WHERE ba.notice_id = pn.id)
        ORDER BY pn.notice_date DESC NULLS LAST
    """)
    summary = {"processed": 0, "bidders_found": 0, "bidders_inserted": 0, "links_created": 0, "errors": 0}
    for row in notices:
        try:
            res = import_bidders_from_notice(row["id"], fetch_detail)
            summary["processed"] += 1
            summary["bidders_found"] += res.get("bidders_found", 0)
            summary["bidders_inserted"] += res.get("bidders_inserted", 0)
            summary["links_created"] += res.get("links_created", 0)
        except Exception:
            summary["errors"] += 1
    return summary


@router.post("/import_awards_all")
def import_awards_all(
    fetch_detail: bool = Query(False),
    country: str | None = Query(None),
    _: dict = Depends(require_auth),
):
    """Batch import bidders for all award notices, optionally limited to one country. Runs in background."""

    def run_batch():
        if country:
            notices = q(
                """SELECT id
                   FROM procurement_notices
                   WHERE (notice_type IN ('Contract Award','Award') OR notice_type ILIKE '%Contract Award%' OR status = 'Awarded')
                     AND country = %s""",
                [country],
            )
        else:
            notices = q(
                "SELECT id FROM procurement_notices WHERE notice_type IN ('Contract Award','Award') OR notice_type ILIKE '%Contract Award%' OR status = 'Awarded'"
            )
        summary = {
            "processed": 0,
            "bidders_found": 0,
            "bidders_inserted": 0,
            "links_created": 0,
            "country": country,
        }
        for n in notices:
            nid = n["id"]
            try:
                res = import_bidders_from_notice(nid, fetch_detail)
                summary["processed"] += 1
                summary["bidders_found"] += res.get("bidders_found", 0)
                summary["bidders_inserted"] += res.get("bidders_inserted", 0)
                summary["links_created"] += res.get("links_created", 0)
            except Exception:
                continue
        try:
            with open("bidders_import_summary.json", "w", encoding="utf-8") as f:
                f.write(json.dumps(summary))
        except Exception:
            pass

    threading.Thread(target=run_batch, daemon=True).start()
    return {"status": "started", "country": country}


@router.post("/import_by_country")
def import_bidders_by_country(
    country: str = Query(..., description="Import bidders only for this country"),
    fetch_detail: bool = Query(False),
    _: dict = Depends(require_auth),
):
    if not country or not country.strip():
        raise HTTPException(status_code=400, detail="country is required")
    return import_awards_all(fetch_detail=fetch_detail, country=country.strip())


@router.post("/import_missing_by_country")
def import_missing_awards_by_country(
    country: str = Query(..., description="Import missing bidder links only for this country"),
    fetch_detail: bool = Query(False),
    _: dict = Depends(require_auth),
):
    normalized_country = (country or "").strip()
    if not normalized_country:
        raise HTTPException(status_code=400, detail="country is required")

    notices = q(
        """
        SELECT pn.id
        FROM procurement_notices pn
        WHERE (pn.notice_type IN ('Contract Award','Award') OR pn.notice_type ILIKE '%Contract Award%' OR pn.status = 'Awarded')
          AND pn.country = %s
          AND NOT EXISTS (SELECT 1 FROM bidder_awards ba WHERE ba.notice_id = pn.id)
        ORDER BY pn.notice_date DESC NULLS LAST
    """,
        [normalized_country],
    )
    summary = {
        "processed": 0,
        "bidders_found": 0,
        "bidders_inserted": 0,
        "links_created": 0,
        "errors": 0,
        "country": normalized_country,
    }
    for row in notices:
        try:
            res = import_bidders_from_notice(row["id"], fetch_detail)
            summary["processed"] += 1
            summary["bidders_found"] += res.get("bidders_found", 0)
            summary["bidders_inserted"] += res.get("bidders_inserted", 0)
            summary["links_created"] += res.get("links_created", 0)
        except Exception:
            summary["errors"] += 1
    return summary


# ── Contact enrichment ────────────────────────────────────────────────────────


@router.post("/enrich", dependencies=[Depends(rate_limited(STRICT_LIMITER))])
def enrich_all_bidders(
    missing_only: bool = Query(True, description="Only enrich bidders missing contact info"),
    limit: int = Query(50, ge=1, le=500),
    _: dict = Depends(require_auth),
):
    if missing_only:
        rows = q(
            """
            SELECT id, name, country FROM bidders
            WHERE (contact_email IS NULL OR TRIM(contact_email) = '')
              AND (contact_phone IS NULL OR TRIM(contact_phone) = '')
            ORDER BY updated_at ASC NULLS FIRST
            LIMIT %s
        """,
            [limit],
        )
    else:
        rows = q(
            """
            SELECT id, name, country FROM bidders
            ORDER BY updated_at ASC NULLS FIRST
            LIMIT %s
        """,
            [limit],
        )

    results = []
    for row in rows:
        try:
            found = search_company_contact(row["name"], row.get("country") or "")
        except Exception as e:
            results.append(
                EnrichResult(status="error", bidder_id=row["id"], bidder_name=row["name"], error=str(e)).dict()
            )
            continue

        if not found:
            results.append(EnrichResult(status="no_results", bidder_id=row["id"], bidder_name=row["name"]).dict())
            continue

        update_fields = {}
        for key in ("contact_email", "contact_phone", "linkedin_url"):
            val = found.get(key)
            if val:
                update_fields[key] = val

        if update_fields:
            sets = ", ".join(f"{k} = %s" for k in update_fields)
            params = list(update_fields.values()) + [row["id"]]
            with db() as conn:
                with conn.cursor() as cur:
                    cur.execute(f"UPDATE bidders SET {sets}, updated_at = NOW() WHERE id = %s", params)
                conn.commit()

        results.append(
            EnrichResult(
                status="ok",
                bidder_id=row["id"],
                bidder_name=row["name"],
                found=found,
                updated_fields=list(update_fields.keys()),
            ).dict()
        )

    enriched_count = sum(1 for r in results if r["status"] == "ok")
    return {
        "total_processed": len(results),
        "enriched": enriched_count,
        "results": results,
    }


@router.post("/{bidder_id}/enrich", dependencies=[Depends(rate_limited(STRICT_LIMITER))])
def enrich_bidder_contact(bidder_id: int, _: dict = Depends(require_auth)):
    rows = q("SELECT * FROM bidders WHERE id = %s", [bidder_id])
    if not rows:
        raise HTTPException(status_code=404, detail="Bidder not found")
    bidder = rows[0]
    company = bidder["name"]
    country = bidder.get("country") or ""

    try:
        found = search_company_contact(company, country)
    except Exception as e:
        return EnrichResult(status="error", bidder_id=bidder_id, bidder_name=company, error=str(e)).dict()

    if not found:
        return EnrichResult(
            status="no_results", bidder_id=bidder_id, bidder_name=company, found={}, updated_fields=[]
        ).dict()

    update_fields = {}
    for key in ("contact_email", "contact_phone", "linkedin_url"):
        val = found.get(key)
        if val:
            update_fields[key] = val

    if update_fields:
        sets = ", ".join(f"{k} = %s" for k in update_fields)
        params = list(update_fields.values()) + [bidder_id]
        with db() as conn:
            with conn.cursor() as cur:
                cur.execute(f"UPDATE bidders SET {sets}, updated_at = NOW() WHERE id = %s", params)
            conn.commit()

    return EnrichResult(
        status="ok", bidder_id=bidder_id, bidder_name=company, found=found, updated_fields=list(update_fields.keys())
    ).dict()


@router.post("/{bidder_id}/enrich_gemini", dependencies=[Depends(rate_limited(STRICT_LIMITER))])
def enrich_bidder_gemini(bidder_id: int, _: dict = Depends(require_auth)):
    # 1. Fetch bidder details
    rows = q("SELECT * FROM bidders WHERE id = %s", [bidder_id])
    if not rows:
        raise HTTPException(status_code=404, detail="Bidder not found")
    bidder = rows[0]

    # 2. Fetch bidder notices for context
    notices = get_bidder_notices(bidder_id)

    # 3. Call Gemini service
    try:
        enriched = classify_and_enrich_with_gemini(
            company=bidder["name"], country=bidder.get("country") or "", bids=notices
        )
    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve)) from ve
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e)) from e

    # 4. Save results to DB
    update_fields = {}
    for key in (
        "category",
        "business_model",
        "core_products",
        "corporate_activities",
        "contact_email",
        "contact_phone",
        "linkedin_url",
    ):
        val = enriched.get(key)
        if val is not None:
            update_fields[key] = val

    if update_fields:
        sets = ", ".join(f"{k} = %s" for k in update_fields)
        params = list(update_fields.values()) + [bidder_id]
        with db() as conn:
            with conn.cursor() as cur:
                cur.execute(f"UPDATE bidders SET {sets}, updated_at = NOW() WHERE id = %s", params)
            conn.commit()

    return {
        "status": "ok",
        "bidder_id": bidder_id,
        "bidder_name": bidder["name"],
        "enriched": enriched,
        "updated_fields": list(update_fields.keys()),
    }


# ── Export ────────────────────────────────────────────────────────────────────


@router.get("/export")
def export_bidders_excel(
    qs: str | None = Query(None),
    search: str | None = Query(None),
    country: str | None = Query(None),
    won_only: bool = Query(False),
    tech_only: bool = Query(False),
    fields: str | None = Query(None),
):
    selected_fields = resolve_bidder_export_fields(fields)
    rows = fetch_bidder_export_rows(search, qs, country, won_only, tech_only)

    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)

    rows_by_country: dict[str, list] = {}
    for row in rows:
        country_key = (row.get("country") or "").strip() or "Unknown"
        rows_by_country.setdefault(country_key, []).append(row)

    used_titles = set()
    for index, country in enumerate(sorted(rows_by_country), 1):
        base_title = _safe_sheet_title(country, f"Country {index}")
        title = base_title
        suffix = 2
        while title.lower() in used_titles:
            tail = f" {suffix}"
            title = f"{base_title[:31 - len(tail)]}{tail}"
            suffix += 1
        used_titles.add(title.lower())

        ws = wb.create_sheet(title)
        _write_bidder_export_sheet(ws, selected_fields, rows_by_country[country])

    if not wb.worksheets:
        ws = wb.create_sheet("Bidders")
        _write_bidder_export_sheet(ws, selected_fields, [])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"Bidders_Custom_{datetime.utcnow().strftime('%Y%m%d_%H%M')}.xlsx"
    from fastapi.responses import StreamingResponse

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export/csv")
def export_bidders_csv(
    qs: str | None = Query(None),
    search: str | None = Query(None),
    country: str | None = Query(None),
    won_only: bool = Query(False),
    tech_only: bool = Query(False),
    fields: str | None = Query(None),
):
    selected_fields = resolve_bidder_export_fields(fields)
    rows = fetch_bidder_export_rows(search, qs, country, won_only, tech_only)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([BIDDER_EXPORT_FIELDS[field][0] for field in selected_fields])
    for row in rows:
        writer.writerow([row.get(BIDDER_EXPORT_FIELDS[field][1]) for field in selected_fields])

    filename = f"Bidders_Custom_{datetime.utcnow().strftime('%Y%m%d_%H%M')}.csv"
    from fastapi.responses import StreamingResponse

    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
