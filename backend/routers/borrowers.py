import csv
import io
import re
from datetime import datetime
from typing import Any

import openpyxl
from db import q
from fastapi import APIRouter, Query
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from services.tech import build_tech_notice_condition

router = APIRouter(prefix="/api/borrowers", tags=["borrowers"])

# ── Export config ─────────────────────────────────────────────────────────────

BORROWER_EXPORT_FIELD_CONFIG = {
    "borrower": ("Institution Name", "borrower", 40),
    "country": ("Country", "country", 18),
    "total_notices": ("Total Notices", "total_notices", 14),
    "ifb_count": ("IFB Notices", "ifb_count", 12),
    "reoi_count": ("REOI Notices", "reoi_count", 12),
    "award_count": ("Award Notices", "award_count", 14),
    "first_notice_date": ("First Notice Date", "first_notice_date", 18),
    "last_notice_date": ("Last Notice Date", "last_notice_date", 18),
    "recent_30d": ("30-Day Activity", "recent_30d", 14),
}

DEFAULT_BORROWER_EXPORT_FIELDS = [
    "borrower",
    "country",
    "total_notices",
    "ifb_count",
    "reoi_count",
    "award_count",
    "first_notice_date",
    "last_notice_date",
    "recent_30d",
]


def resolve_borrower_export_fields(fields: str | None) -> list[str]:
    selected_fields = [f.strip() for f in (fields or "").split(",") if f.strip()]
    if not selected_fields:
        return DEFAULT_BORROWER_EXPORT_FIELDS.copy()
    return [field for field in selected_fields if field in BORROWER_EXPORT_FIELD_CONFIG]


def fetch_borrower_export_rows(search, country, notice_type):
    filters = ["borrower IS NOT NULL AND TRIM(borrower) <> ''"]
    params = []

    if search:
        filters.append("borrower ILIKE %s")
        params.append(f"%{search}%")
    if country:
        filters.append("country = %s")
        params.append(country)
    if notice_type:
        filters.append("""(
            notice_type = %s OR
            notice_type ILIKE %s OR
            notice_type ILIKE %s
        )""")
        full = {
            "IFB": "%Invitation for Bids%",
            "REOI": "%Expression of Interest%",
            "Contract Award": "%Contract Award%",
            "Award": "%Contract Award%",
        }
        params.append(notice_type)
        params.append(full.get(notice_type, f"%{notice_type}%"))
        params.append(f"%{notice_type}%")

    where = " AND ".join(filters)

    return q(
        f"""
        SELECT
            borrower,
            country,
            COUNT(*)                                                     AS total_notices,
            COUNT(*) FILTER (WHERE notice_type = 'IFB')                  AS ifb_count,
            COUNT(*) FILTER (WHERE notice_type = 'REOI')                 AS reoi_count,
            COUNT(*) FILTER (WHERE notice_type IN ('Award','Contract Award')) AS award_count,
            MIN(notice_date)::text                                       AS first_notice_date,
            MAX(notice_date)::text                                       AS last_notice_date,
            COUNT(*) FILTER (WHERE notice_date >= CURRENT_DATE - INTERVAL '30 days') AS recent_30d
        FROM procurement_notices
        WHERE {where}
        GROUP BY borrower, country
        ORDER BY total_notices DESC, borrower ASC
    """,
        params,
    )


# ── Endpoints ─────────────────────────────────────────────────────────────────


@router.get("")
def list_borrowers(
    search: str | None = Query(None, description="Search by borrower name"),
    country: str | None = Query(None, description="Filter by country"),
    notice_type: str | None = Query(None, description="Filter by notice type"),
    page: int = Query(1, ge=1),
    page_size: int = Query(25, le=100),
):
    filters = ["borrower IS NOT NULL AND TRIM(borrower) <> ''"]
    params = []

    if search:
        filters.append("borrower ILIKE %s")
        params.append(f"%{search}%")
    if country:
        filters.append("country = %s")
        params.append(country)
    if notice_type:
        filters.append("""(
            notice_type = %s OR
            notice_type ILIKE %s OR
            notice_type ILIKE %s
        )""")
        full = {
            "IFB": "%Invitation for Bids%",
            "REOI": "%Expression of Interest%",
            "Contract Award": "%Contract Award%",
            "Award": "%Contract Award%",
        }
        params.append(notice_type)
        params.append(full.get(notice_type, f"%{notice_type}%"))
        params.append(f"%{notice_type}%")

    where = " AND ".join(filters)

    offset = (page - 1) * page_size

    rows = q(
        f"""
        SELECT
            borrower,
            country,
            COUNT(*)                                                     AS total_notices,
            COUNT(*) FILTER (WHERE notice_type = 'IFB')                  AS ifb_count,
            COUNT(*) FILTER (WHERE notice_type = 'REOI')                 AS reoi_count,
            COUNT(*) FILTER (WHERE notice_type IN ('Award','Contract Award')) AS award_count,
            MIN(notice_date)::text                                       AS first_notice_date,
            MAX(notice_date)::text                                       AS last_notice_date,
            COUNT(*) FILTER (WHERE notice_date >= CURRENT_DATE - INTERVAL '30 days') AS recent_30d
        FROM procurement_notices
        WHERE {where}
        GROUP BY borrower, country
        ORDER BY total_notices DESC, borrower ASC
        LIMIT %s OFFSET %s
    """,
        params + [page_size, offset],
    )

    total = len(rows)
    if total == page_size:
        total = q(
            f"""
            SELECT COUNT(*) AS cnt FROM (
                SELECT 1 FROM procurement_notices
                WHERE {where}
                GROUP BY borrower, country
            ) sub
        """,
            params,
        )[0]["cnt"]

    output = []
    for row in rows:
        item = dict(row)
        item["total_notices"] = int(item["total_notices"])
        item["ifb_count"] = int(item["ifb_count"])
        item["reoi_count"] = int(item["reoi_count"])
        item["award_count"] = int(item["award_count"])
        item["recent_30d"] = int(item["recent_30d"])
        output.append(item)

    available_countries = [
        r["country"]
        for r in q("""
        SELECT DISTINCT country FROM procurement_notices
        WHERE country IS NOT NULL AND TRIM(country) <> ''
        ORDER BY country
    """)
    ]

    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "data": output,
        "available_countries": available_countries,
    }


@router.get("/export")
def export_borrowers(
    search: str | None = Query(None),
    country: str | None = Query(None),
    notice_type: str | None = Query(None),
    fields: str | None = Query(None),
    format: str = Query("xlsx", regex="^(xlsx|csv)$"),
):
    selected_fields = resolve_borrower_export_fields(fields)
    rows = fetch_borrower_export_rows(search, country, notice_type)

    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        headers = [BORROWER_EXPORT_FIELD_CONFIG[f][0] for f in selected_fields]
        writer.writerow(headers)
        for row in rows:
            writer.writerow([row.get(BORROWER_EXPORT_FIELD_CONFIG[f][1]) for f in selected_fields])
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M")
        return StreamingResponse(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="WB_Borrowers_{timestamp}.csv"'},
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Borrowers"

    header_fill = PatternFill("solid", fgColor="2C3E50")
    header_font = Font(bold=True, color="FFFFFF", size=12, name="Calibri")
    header_border = Border(
        left=Side(style="thin", color="34495E"),
        right=Side(style="thin", color="34495E"),
        top=Side(style="medium", color="3498DB"),
        bottom=Side(style="medium", color="3498DB"),
    )
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_font = Font(size=11, name="Calibri")
    border_all = Border(
        left=Side(style="thin", color="BDC3C7"),
        right=Side(style="thin", color="BDC3C7"),
        top=Side(style="thin", color="BDC3C7"),
        bottom=Side(style="thin", color="BDC3C7"),
    )
    center_align = Alignment(horizontal="center", vertical="center")

    COLUMNS = [BORROWER_EXPORT_FIELD_CONFIG[f] for f in selected_fields]

    for col_idx, (label, _, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = header_border
        cell.alignment = header_align

    ws.row_dimensions[1].height = 35
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    ws.sheet_view.zoomScale = 85

    for row_idx, row in enumerate(rows, 2):
        ws.row_dimensions[row_idx].height = 22
        for col_idx, (_, key, _) in enumerate(COLUMNS, 1):
            val = row.get(key)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = border_all
            cell.alignment = (
                center_align
                if key in ("total_notices", "ifb_count", "reoi_count", "award_count", "recent_30d")
                else Alignment(horizontal="left", vertical="center")
            )
            if key in ("first_notice_date", "last_notice_date") and val:
                try:
                    if isinstance(val, str) and len(val) >= 10:
                        cell.value = datetime.strptime(val[:10], "%Y-%m-%d").date()
                        cell.alignment = center_align
                except Exception:
                    pass

    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="WB_Borrowers_{timestamp}.xlsx"'},
    )


@router.get("/export/qualified")
def export_qualified_borrowers(
    min_awards: int = Query(3, ge=1),
    tech_only: bool = Query(False, description="Only tech-related awarded notices"),
):
    tech_notice_filter, tech_notice_params = build_tech_notice_condition("pn")
    tech_where = f"AND {tech_notice_filter}" if tech_only else ""
    tech_params = tech_notice_params if tech_only else []

    notices = q(
        f"""
        SELECT DISTINCT
            pn.borrower           AS host_institution,
            pn.country,
            pn.project_id,
            COALESCE(NULLIF(pn.borrower_bid_reference, ''), NULLIF(pn.notice_no, ''), pn.id) AS reference_number,
            pn.title,
            pn.notice_date::text,
            pn.notice_type,
            pn.status,
            pn.contract_amount,
            pn.currency,
            pn.url,
            pn.description
        FROM procurement_notices pn
        WHERE pn.notice_type IN ('Award', 'Contract Award')
          AND pn.borrower IS NOT NULL
          AND TRIM(pn.borrower) <> ''
          {tech_where}
          AND (pn.borrower, pn.country) IN (
              SELECT borrower, country
              FROM procurement_notices
              WHERE notice_type IN ('Award', 'Contract Award')
                AND borrower IS NOT NULL
                AND TRIM(borrower) <> ''
              GROUP BY borrower, country
              HAVING COUNT(*) >= %s
          )
        ORDER BY country, host_institution, notice_date DESC NULLS LAST
    """,
        tech_params + [min_awards],
    )

    if not notices:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "No Data"
        ws.cell(row=1, column=1, value="No qualifying institutions found")
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M")
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="Qualified_Institutions_{timestamp}.xlsx"'},
        )

    countries_order = []
    rows_by_country: dict[str, list[dict[str, Any]]] = {}
    for row in notices:
        ref = row.get("reference_number") or ""
        if ref.startswith("OP") or ref.startswith("PN") or ref.startswith("00"):
            desc = row.get("description") or ""
            m = re.search(r"(?:Bid/Contract\s+)?Reference\s+No\.?\s*[:\s]\s*\n?\s*(.+)", desc)
            if m:
                extracted = m.group(1).strip()
                if extracted and not extracted.startswith("0") and len(extracted) < 60:
                    row["reference_number"] = extracted
        c = row.get("country") or "Unknown"
        rows_by_country.setdefault(c, []).append(row)
        if c not in countries_order:
            countries_order.append(c)

    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    header_fill = PatternFill("solid", fgColor="1B5E20")
    header_font = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
    header_border = Border(
        left=Side(style="thin", color="2E7D32"),
        right=Side(style="thin", color="2E7D32"),
        top=Side(style="medium", color="43A047"),
        bottom=Side(style="medium", color="43A047"),
    )
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_font = Font(size=10, name="Calibri")
    border_all = Border(
        left=Side(style="thin", color="C8E6C9"),
        right=Side(style="thin", color="C8E6C9"),
        top=Side(style="thin", color="C8E6C9"),
        bottom=Side(style="thin", color="C8E6C9"),
    )
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    COLUMNS = [
        ("Host Institution", "host_institution", 45),
        ("Tender Name", "title", 55),
        ("Project ID", "project_id", 22),
        ("Reference No.", "reference_number", 30),
    ]

    used_titles = set()
    for country in countries_order:
        country_rows = rows_by_country[country]
        base_title = _safe_sheet_title(country, country)
        title = base_title
        suffix = 2
        while title.lower() in used_titles:
            tail = f" {suffix}"
            title = f"{base_title[:31 - len(tail)]}{tail}"
            suffix += 1
        used_titles.add(title.lower())

        ws = wb.create_sheet(title)

        ws.row_dimensions[1].height = 32
        for col_idx, (label, _, width) in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = header_border
            cell.alignment = header_align
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
        ws.sheet_view.zoomScale = 90

        for row_idx, row in enumerate(country_rows, 2):
            ws.row_dimensions[row_idx].height = 22
            for col_idx, (_, key, _) in enumerate(COLUMNS, 1):
                val = row.get(key)
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = border_all
                cell.alignment = left_align

    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M")
    tech_tag = "_Tech" if tech_only else ""
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="Qualified_Institutions{tech_tag}_{timestamp}.xlsx"'},
    )


def _safe_sheet_title(title: str, fallback: str) -> str:
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "", (title or "").strip())
    return (cleaned or fallback)[:31]
