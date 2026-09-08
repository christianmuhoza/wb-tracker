import csv
import io
import re
from datetime import datetime
from typing import Any

import openpyxl
from bs4 import BeautifulSoup
from config import DEFAULT_COUNTRIES
from db import q
from fastapi import APIRouter, Query
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from services.bidder_extraction import (
    extract_awarded_bidders,
    extract_bidders_from_description,
    extract_bidders_list,
    parse_notice_bidder_details,
)
from services.tech import classify_notice_tech

router = APIRouter(prefix="/api/export", tags=["export"])

DEFAULT_COUNTRIES = list(DEFAULT_COUNTRIES)

# ── Export config ─────────────────────────────────────────────────────────────

EXPORT_FIELD_CONFIG = {
    "country": ("Country", "country", 15),
    "notice_type": ("Notice Type", "notice_type", 16),
    "status": ("Status", "status", 12),
    "title": ("Opportunity Title", "title", 50),
    "project_id": ("Project ID", "project_id", 18),
    "project_name": ("Project Name", "project_name", 40),
    "procurement_method": ("Procurement Method", "procurement_method", 25),
    "borrower": ("Host Institution", "borrower", 35),
    "notice_date": ("Notice Date", "notice_date", 15),
    "awarded_date": ("Awarded Date", "awarded_date", 15),
    "submission_date": ("Submission Deadline", "submission_date", 20),
    "contract_amount": ("Contract Amount", "contract_amount", 20),
    "currency": ("Currency", "currency", 10),
    "contact_email": ("Contact Email", "contact_email", 30),
    "url": ("World Bank Link", "url", 25),
    "bidders": ("Bidders", "bidders", 40),
    "is_tech": ("Tech Opportunity", "is_tech", 16),
    "tech_category": ("Tech Category", "tech_category", 28),
    "overview": ("Overview", "overview", 60),
    "requirements": ("Requirements", "requirements", 60),
    "description": ("Description", "description", 60),
    "fetched_at": ("Fetched Date", "fetched_at", 18),
    "opportunity_id": ("Opportunity ID", "opportunity_id", 20),
}

CUSTOM_BIDDER_EXPORT_FIELD_CONFIG = {
    "country": ("Country", "country", 15),
    "notice_type": ("Notice Type", "notice_type", 16),
    "status": ("Status", "status", 12),
    "title": ("Opportunity Title", "title", 50),
    "project_id": ("Project ID", "project_id", 18),
    "project_name": ("Project Name", "project_name", 40),
    "procurement_method": ("Procurement Method", "procurement_method", 25),
    "borrower": ("Host Institution", "borrower", 35),
    "notice_date": ("Notice Date", "notice_date", 15),
    "submission_date": ("Submission Deadline", "submission_date", 20),
    "is_tech": ("Tech Opportunity", "is_tech", 16),
    "tech_category": ("Tech Category", "tech_category", 28),
    "overview": ("Overview", "overview", 60),
    "requirements": ("Requirements", "requirements", 60),
    "bidder_name": ("Bidder Name", "bidder_name", 40),
    "bidder_country": ("Bidder Country", "bidder_country", 20),
    "bidder_status": ("Bidder Status", "bidder_status", 18),
    "won": ("Won", "won", 10),
    "bid_price_at_opening": ("Bid Price at Opening", "bid_price_at_opening", 20),
    "opening_currency": ("Opening Bid Currency", "opening_currency", 16),
    "evaluated_bid_price": ("Evaluated Bid Price", "evaluated_bid_price", 20),
    "evaluated_bid_currency": ("Evaluated Bid Currency", "evaluated_bid_currency", 18),
    "final_evaluation_price": ("Final Evaluation Price", "final_evaluation_price", 20),
    "final_evaluation_currency": ("Final Evaluation Currency", "final_evaluation_currency", 18),
    "winner_contract_amount": ("Contract Amount (Winner Only)", "winner_contract_amount", 24),
    "winner_contract_currency": ("Contract Currency", "winner_contract_currency", 16),
    "contact_email": ("Contact Email", "contact_email", 30),
    "url": ("World Bank Link", "url", 25),
    "opportunity_id": ("Opportunity ID", "opportunity_id", 20),
}

EXPORT_FIELD_MAPPING = {
    "country": "country",
    "notice_type": "notice_type",
    "status": "status",
    "title": "title",
    "project_id": "project_id",
    "project_name": "project_name",
    "procurement_method": "procurement_method",
    "borrower": "borrower",
    "notice_date": "notice_date",
    "awarded_date": "(SELECT MAX(ba.award_date) FROM bidder_awards ba WHERE ba.notice_id = procurement_notices.id) AS awarded_date",
    "submission_date": "COALESCE(submission_date, submission_deadline::date)::text AS submission_date",
    "contract_amount": "contract_amount",
    "currency": "currency",
    "contact_email": "contact_email",
    "url": "url",
    "description": "description",
    "fetched_at": "fetched_at",
    "opportunity_id": "id as opportunity_id",
    "bidders": "description",
    "overview": "description",
    "requirements": "description",
    "is_tech": "description",
    "tech_category": "description",
}

DEFAULT_EXPORT_FIELDS = [
    "country",
    "notice_type",
    "status",
    "title",
    "project_id",
    "project_name",
    "procurement_method",
    "borrower",
    "notice_date",
    "awarded_date",
    "submission_date",
    "contract_amount",
    "currency",
    "contact_email",
    "url",
]

BORROWER_EXPORT_FIELD_CONFIG = {
    "borrower": ("Institution Name", "borrower", 40),
    "country": ("Country", "country", 18),
    "total_notices": ("Total Notices", "total_notices", 14),
    "ifb_count": ("IFB Notices", "ifb_count", 12),
    "reoi_count": ("REOI Notices", "reoi_count", 12),
    "award_count": ("Award Notices", "award_count", 14),
    "first_notice_date": ("First Notice Date", "first_notice_date", 18),
    "last_notice_date": ("Last Notice Date", "last_notice_date", 18),
    "recent_30d": ("30-Day Activity", "recent_30d", 14),
}

DEFAULT_BORROWER_EXPORT_FIELDS = [
    "borrower",
    "country",
    "total_notices",
    "ifb_count",
    "reoi_count",
    "award_count",
    "first_notice_date",
    "last_notice_date",
    "recent_30d",
]

DEFAULT_CUSTOM_BIDDER_EXPORT_FIELDS = [
    "country",
    "notice_type",
    "title",
    "project_id",
    "project_name",
    "borrower",
    "notice_date",
    "bidder_name",
    "bidder_country",
    "bidder_status",
    "won",
    "evaluated_bid_price",
    "evaluated_bid_currency",
    "winner_contract_amount",
    "winner_contract_currency",
    "url",
]

BASE_TEMPLATE_EXPORT_COLUMNS = [
    ("CONTRACTOR", "contractor", 20.6640625),
    ("CONTRACTOR'S COUNTRY", "contractor_country", 35.77734375),
    ("TENDER", "tender", 19.88671875),
    ("AWARD DATE", "award_date", 31.5546875),
    ("AWARD AMOUNT", "award_amount", 41.109375),
    ("CONTRACTOR REFERENCE NUMBER", "contractor_reference_number", 59.6640625),
    ("CONTACT", "contact", 31.109375),
]


def resolve_export_fields(fields: str | None) -> list[str]:
    selected_fields = [f.strip() for f in (fields or "").split(",") if f.strip()]
    if not selected_fields:
        return DEFAULT_EXPORT_FIELDS.copy()
    return [field for field in selected_fields if field in EXPORT_FIELD_CONFIG]


def resolve_custom_bidder_export_fields(fields: str | None) -> list[str]:
    selected_fields = [f.strip() for f in (fields or "").split(",") if f.strip()]
    if not selected_fields:
        return DEFAULT_CUSTOM_BIDDER_EXPORT_FIELDS.copy()
    return [field for field in selected_fields if field in CUSTOM_BIDDER_EXPORT_FIELD_CONFIG]


# ── Shared filter builder ─────────────────────────────────────────────────────


def build_where(country, notice_type, status, from_date, to_date, search):
    filters = ["1=1"]
    params = []

    if country:
        filters.append("country = %s")
        params.append(country)
    if notice_type:
        filters.append("""(
            notice_type = %s OR
            notice_type ILIKE %s OR
            notice_type ILIKE %s
        )""")
        full = {
            "IFB": "%Invitation for Bids%",
            "REOI": "%Expression of Interest%",
            "Contract Award": "%Contract Award%",
            "Award": "%Contract Award%",
        }
        params.append(notice_type)
        params.append(full.get(notice_type, f"%{notice_type}%"))
        params.append(f"%{notice_type}%")
    if status:
        filters.append("status ILIKE %s")
        params.append(status)
    if from_date:
        filters.append("notice_date >= %s")
        params.append(from_date)
    if to_date:
        filters.append("notice_date <= %s")
        params.append(to_date)
    if search:
        filters.append("(title ILIKE %s OR description ILIKE %s OR project_name ILIKE %s)")
        like = f"%{search}%"
        params.extend([like, like, like])

    return " AND ".join(filters), params


# ── Row fetchers ──────────────────────────────────────────────────────────────


def fetch_export_rows(where: str, params, selected_fields: list[str]):
    select_fields = []
    for field in selected_fields:
        mapped = EXPORT_FIELD_MAPPING.get(field)
        if mapped:
            select_fields.append(mapped)
    if "bidders" in selected_fields and "description" not in select_fields:
        select_fields.append("description")
    if any(field in selected_fields for field in ("is_tech", "tech_category")):
        for extra in ("title", "project_name", "procurement_method", "borrower_bid_reference", "description"):
            if extra not in select_fields:
                select_fields.append(extra)
    if not select_fields:
        select_fields = ["country", "title", "notice_date"]

    seen = set()
    ordered_select = []
    for field in select_fields:
        if field not in seen:
            seen.add(field)
            ordered_select.append(field)

    return q(
        f"""
        SELECT {', '.join(ordered_select)}
        FROM procurement_notices
        WHERE {where}
        ORDER BY notice_date DESC
    """,
        params or (),
    )


def fetch_custom_bidder_export_rows(where: str, params, selected_fields: list[str]):
    notice_rows = q(
        f"""
        SELECT
            country,
            notice_type,
            status,
            title,
            project_id,
            project_name,
            procurement_method,
            borrower,
            borrower_bid_reference,
            notice_no,
            notice_date::text,
            COALESCE(submission_date, submission_deadline::date)::text AS submission_date,
            contract_amount,
            currency,
            contact_email,
            url,
            id AS opportunity_id,
            description
        FROM procurement_notices
        WHERE {where}
        ORDER BY notice_date DESC
    """,
        params or (),
    )

    bidder_rows: list[dict[str, Any]] = []
    base_notice_fields = (
        "country",
        "notice_type",
        "status",
        "title",
        "project_id",
        "project_name",
        "procurement_method",
        "borrower",
        "notice_date",
        "submission_date",
        "contact_email",
        "url",
        "opportunity_id",
        "borrower_bid_reference",
        "notice_no",
        "description",
    )

    for notice in notice_rows:
        desc = notice.get("description") or ""
        bidder_details = parse_notice_bidder_details(desc)

        if bidder_details:
            for detail in bidder_details:
                won = detail.get("section") == "awarded"
                row = {field: notice.get(field) for field in base_notice_fields}
                row.update(
                    {
                        "bidder_name": detail.get("name"),
                        "bidder_country": detail.get("country"),
                        "bidder_status": detail.get("role"),
                        "won": "Yes" if won else "No",
                        "bid_price_at_opening": detail.get("opening_amount"),
                        "opening_currency": detail.get("opening_currency"),
                        "evaluated_bid_price": detail.get("evaluated_amount"),
                        "evaluated_bid_currency": detail.get("evaluated_currency"),
                        "final_evaluation_price": detail.get("final_amount"),
                        "final_evaluation_currency": detail.get("final_currency"),
                        "winner_contract_amount": (
                            detail.get("signed_amount")
                            if won and detail.get("signed_amount") is not None
                            else notice.get("contract_amount")
                            if won
                            else None
                        ),
                        "winner_contract_currency": (
                            (detail.get("signed_currency") or detail.get("currency") or notice.get("currency"))
                            if won
                            else None
                        ),
                    }
                )
                bidder_rows.append(row)
            continue

        all_bidders = extract_bidders_list(desc)
        awarded_names = {name.lower() for name in extract_awarded_bidders(desc)}
        for bidder_name in all_bidders:
            won = bidder_name.lower() in awarded_names
            bidder_rows.append(
                {
                    "country": notice.get("country"),
                    "notice_type": notice.get("notice_type"),
                    "status": notice.get("status"),
                    "title": notice.get("title"),
                    "project_id": notice.get("project_id"),
                    "project_name": notice.get("project_name"),
                    "procurement_method": notice.get("procurement_method"),
                    "borrower": notice.get("borrower"),
                    "borrower_bid_reference": notice.get("borrower_bid_reference"),
                    "notice_no": notice.get("notice_no"),
                    "notice_date": notice.get("notice_date"),
                    "submission_date": notice.get("submission_date"),
                    "bidder_name": bidder_name,
                    "bidder_country": None,
                    "bidder_status": "Awarded" if won else "Bidder",
                    "won": "Yes" if won else "No",
                    "bid_price_at_opening": None,
                    "opening_currency": None,
                    "evaluated_bid_price": None,
                    "evaluated_bid_currency": None,
                    "final_evaluation_price": None,
                    "final_evaluation_currency": None,
                    "winner_contract_amount": notice.get("contract_amount") if won else None,
                    "winner_contract_currency": notice.get("currency") if won else None,
                    "contact_email": notice.get("contact_email"),
                    "url": notice.get("url"),
                    "opportunity_id": notice.get("opportunity_id"),
                    "description": notice.get("description"),
                }
            )

    return bidder_rows


def fetch_template_export_rows(where: str, params):
    linked_rows = q(
        f"""
        WITH filtered_notices AS (
            SELECT *
            FROM procurement_notices
            WHERE {where}
        )
        SELECT
            pn.country,
            b.name AS contractor,
            b.country AS contractor_country,
            pn.title AS tender,
            COALESCE(ba.award_date, pn.notice_date)::text AS award_date,
            COALESCE(pn.submission_date, pn.submission_deadline::date)::text AS deadline_date,
            COALESCE(ba.award_amount, pn.contract_amount) AS award_amount,
            COALESCE(NULLIF(ba.currency, ''), NULLIF(pn.currency, '')) AS award_currency,
            COALESCE(NULLIF(pn.borrower_bid_reference, ''), NULLIF(pn.notice_no, ''), pn.id) AS contractor_reference_number,
            COALESCE(NULLIF(b.contact_email, ''), NULLIF(pn.contact_email, '')) AS contact
        FROM filtered_notices pn
        JOIN bidder_awards ba ON ba.notice_id = pn.id
        JOIN bidders b ON b.id = ba.bidder_id
        WHERE ba.won IS TRUE
        ORDER BY pn.country ASC, COALESCE(ba.award_date, pn.notice_date) DESC NULLS LAST, b.name ASC
    """,
        params or (),
    )

    if linked_rows:
        return [
            {
                "country": row.get("country"),
                "contractor": row.get("contractor"),
                "contractor_country": row.get("contractor_country"),
                "tender": row.get("tender"),
                "award_date": row.get("award_date"),
                "deadline_date": row.get("deadline_date"),
                "award_amount": _format_template_amount(row.get("award_amount"), row.get("award_currency")),
                "contractor_reference_number": row.get("contractor_reference_number"),
                "contact": row.get("contact"),
            }
            for row in linked_rows
        ]

    rows = fetch_custom_bidder_export_rows(where, params, DEFAULT_CUSTOM_BIDDER_EXPORT_FIELDS)
    output = []
    for row in rows:
        if row.get("won") != "Yes":
            continue
        amount = row.get("winner_contract_amount")
        currency = row.get("winner_contract_currency")
        output.append(
            {
                "country": row.get("country"),
                "contractor": row.get("bidder_name"),
                "contractor_country": row.get("bidder_country"),
                "tender": row.get("title"),
                "award_date": row.get("notice_date"),
                "deadline_date": row.get("submission_date"),
                "award_amount": _format_template_amount(amount, currency),
                "contractor_reference_number": (
                    row.get("borrower_bid_reference") or row.get("notice_no") or row.get("opportunity_id")
                ),
                "contact": row.get("contact_email"),
            }
        )
    return output


# ── Helpers ───────────────────────────────────────────────────────────────────


def _format_template_amount(amount, currency):
    if amount is None:
        return None
    currency_text = (currency or "").strip()
    try:
        number = f"{float(amount):,.2f}"
    except Exception:
        number = str(amount)
    symbols = {
        "USD": "$",
        "US$": "$",
        "EUR": "\u20ac",
        "GBP": "\u00a3",
    }
    prefix = symbols.get(currency_text.upper(), f"{currency_text} " if currency_text else "")
    return f"{prefix}{number}"


def get_template_export_columns(include_deadline: bool = False):
    columns = BASE_TEMPLATE_EXPORT_COLUMNS.copy()
    if include_deadline:
        columns.insert(4, ("DEADLINE DATE", "deadline_date", 20))
    return columns + [(None, f"blank_{i}", 13) for i in range(len(columns) + 1, 26)]


def _clean_export_text(text: str | None) -> str:
    if not text:
        return ""
    cleaned = BeautifulSoup(str(text), "html.parser").get_text(" ")
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned.strip()


def extract_notice_overview(text: str | None) -> str | None:
    cleaned = _clean_export_text(text)
    if not cleaned:
        return None

    overview_patterns = [
        r"(?:Scope of Contract|Scope of Work|Description|Assignment Title|Project Description)\s*:?\s*(.+?)(?=\s+(?:Loan/Credit/TF Info|Bid/Contract Reference No|Procurement Method|Awarded Bidder|Evaluated Bidder|Qualification|Eligibility|Requirements?)\s*:|$)",
        r"(?:The objective of|The project (?:will|is)|This assignment (?:will|is)|The consulting services include)\s+(.+?)(?=\s+(?:Qualification|Eligibility|Requirements?|Interested consultants|The attention of)\b|$)",
    ]
    for pattern in overview_patterns:
        match = re.search(pattern, cleaned, flags=re.IGNORECASE)
        if match:
            value = match.group(1).strip(" :-")
            if value:
                return value[:1000]
    return cleaned[:1000]


def extract_notice_requirements(text: str | None) -> str | None:
    cleaned = _clean_export_text(text)
    if not cleaned:
        return None

    section_patterns = [
        r"((?:Qualification|Eligibility|Requirements?|Minimum Qualification|Selection Criteria|Evaluation Criteria|Shortlisting Criteria|Interested consultants)[^:]{0,80}:?\s+.+?)(?=\s+(?:Submission|Deadline|Expressions of Interest|Further information|Contact|Awarded Bidder|Evaluated Bidder|Scope of Contract)\b|$)",
        r"((?:The attention of interested Consultants|Interested consultants should provide|Consultants may associate|A Consultant will be selected).+?)(?=\s+(?:Further information|Expressions of Interest|Submission|Deadline|Contact)\b|$)",
    ]
    sections = []
    for pattern in section_patterns:
        for match in re.finditer(pattern, cleaned, flags=re.IGNORECASE):
            value = match.group(1).strip(" :-")
            if value and value.lower() not in {item.lower() for item in sections}:
                sections.append(value[:1000])
    if sections:
        return " | ".join(sections)[:1500]

    hints = []
    keyword_hints = [
        ("qualification", "Professional qualifications or certifications"),
        ("experience", "Relevant experience on similar assignments"),
        ("financial", "Financial capacity"),
        ("technical", "Technical capability"),
        ("equipment", "Required equipment or tools"),
        ("license", "Valid licenses or registration"),
    ]
    lower = cleaned.lower()
    for keyword, label in keyword_hints:
        if keyword in lower:
            hints.append(label)
    return "; ".join(hints) if hints else None


def _safe_sheet_title(title: str, fallback: str) -> str:
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "", (title or "").strip())
    return (cleaned or fallback)[:31]


def populate_template_sheet(ws, rows, include_deadline: bool = False):
    columns = get_template_export_columns(include_deadline)
    ws.row_dimensions[1].height = 30.6

    for col_idx, (label, _key, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        cell = ws.cell(row=1, column=col_idx, value=label)
        if label:
            cell.font = Font(color="FFFFFF", size=11, name="Calibri")
            cell.fill = PatternFill("solid", fgColor="356854")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, row in enumerate(rows, 2):
        for col_idx, (_, key, _) in enumerate(columns, 1):
            value = row.get(key)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(size=11, name="Calibri")
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            if key in ("award_date", "deadline_date") and value:
                try:
                    if isinstance(value, str) and len(value) >= 10:
                        cell.value = datetime.strptime(value[:10], "%Y-%m-%d").date()
                except Exception:
                    pass
            elif key == "award_amount" and value is not None:
                cell.alignment = Alignment(horizontal="right", vertical="center")


def build_template_workbook(rows, include_deadline: bool = False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    populate_template_sheet(ws, rows, include_deadline)
    return wb


def build_country_template_workbook(rows, countries, include_deadline: bool = False):
    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)

    used_titles = set()
    rows_by_country: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        rows_by_country.setdefault(row.get("country") or "", []).append(row)

    for index, country in enumerate(countries, 1):
        base_title = _safe_sheet_title(country, f"Country {index}")
        title = base_title
        suffix = 2
        while title.lower() in used_titles:
            tail = f" {suffix}"
            title = f"{base_title[:31 - len(tail)]}{tail}"
            suffix += 1
        used_titles.add(title.lower())

        ws = wb.create_sheet(title)
        populate_template_sheet(ws, rows_by_country.get(country, []), include_deadline)

    if not wb.worksheets:
        ws = wb.create_sheet("Sheet1")
        populate_template_sheet(ws, rows, include_deadline)

    return wb


# ── Endpoints ─────────────────────────────────────────────────────────────────


@router.get("/")
def export_excel(
    country: str | None = None,
    notice_type: str | None = None,
    status: str | None = None,
    search: str | None = None,
    from_date: str | None = None,
    to_date: str | None = None,
    type: str | None = "normal",
    fields: str | None = None,
    include_deadline: bool = Query(False),
):
    where, params = build_where(
        country=country, notice_type=notice_type, status=status, search=search, from_date=from_date, to_date=to_date
    )

    is_template_export = type in ("template", "contractor_template")
    is_country_template_export = type in ("template_by_country", "contractor_template_by_country")
    is_huzalink_export = type == "huzalink"
    is_custom_bidder_export = bool(fields) and not is_huzalink_export

    if is_template_export or is_country_template_export:
        rows = fetch_template_export_rows(where, params)
        if is_country_template_export:
            configured_rows = q("SELECT name FROM target_countries ORDER BY name")
            countries = [row["name"] for row in configured_rows] or DEFAULT_COUNTRIES
            wb = build_country_template_workbook(rows, countries, include_deadline)
        else:
            wb = build_template_workbook(rows, include_deadline)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M")
        deadline_part = "_With_Deadline" if include_deadline else ""
        filename = (
            f"Contractor_Template_By_Country{deadline_part}_{timestamp}.xlsx"
            if is_country_template_export
            else f"Contractor_Template_Export{deadline_part}_{timestamp}.xlsx"
        )
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    elif is_huzalink_export:
        selected_fields = resolve_export_fields(fields)
        rows = fetch_export_rows(where, params, selected_fields)
    elif is_custom_bidder_export:
        selected_fields = resolve_custom_bidder_export_fields(fields)
        rows = fetch_custom_bidder_export_rows(where, params, selected_fields)
    else:
        selected_fields = DEFAULT_EXPORT_FIELDS.copy()
        rows = fetch_export_rows(where, params, DEFAULT_EXPORT_FIELDS + ["description"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Procurement Notices"

    header_fill = PatternFill("solid", fgColor="2C3E50")
    header_font = Font(bold=True, color="FFFFFF", size=12, name="Calibri")
    header_border = Border(
        left=Side(style="thin", color="34495E"),
        right=Side(style="thin", color="34495E"),
        top=Side(style="medium", color="3498DB"),
        bottom=Side(style="medium", color="3498DB"),
    )
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_font = Font(size=11, name="Calibri")
    border_all = Border(
        left=Side(style="thin", color="BDC3C7"),
        right=Side(style="thin", color="BDC3C7"),
        top=Side(style="thin", color="BDC3C7"),
        bottom=Side(style="thin", color="BDC3C7"),
    )
    link_font = Font(color="3498DB", underline="single", size=11, name="Calibri")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    if is_huzalink_export:
        COLUMNS = [EXPORT_FIELD_CONFIG[field] for field in selected_fields if field in EXPORT_FIELD_CONFIG]
        ws.title = "Huzalink Export - Custom Fields"
    elif is_custom_bidder_export:
        COLUMNS = [
            CUSTOM_BIDDER_EXPORT_FIELD_CONFIG[field]
            for field in selected_fields
            if field in CUSTOM_BIDDER_EXPORT_FIELD_CONFIG
        ]
        ws.title = "Custom Bidder Export"
    else:
        COLUMNS = [EXPORT_FIELD_CONFIG[field] for field in DEFAULT_EXPORT_FIELDS] + [
            EXPORT_FIELD_CONFIG["overview"],
            EXPORT_FIELD_CONFIG["requirements"],
            ("Description", "description", 80),
        ]
        ws.title = "Procurement Notices"

    for col_idx, (label, _, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = header_border
        cell.alignment = header_align

    ws.row_dimensions[1].height = 35
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    ws.sheet_view.zoomScale = 85

    for row_idx, row in enumerate(rows, 2):
        ws.row_dimensions[row_idx].height = 25
        for col_idx, (_, key, _) in enumerate(COLUMNS, 1):
            if key == "bidders":
                desc = row.get("description")
                val = extract_bidders_from_description(desc) if desc else None
            elif key == "overview":
                val = extract_notice_overview(row.get("description"))
            elif key == "requirements":
                val = extract_notice_requirements(row.get("description"))
            elif key == "is_tech":
                val = "Yes" if classify_notice_tech(row).get("is_tech") else "No"
            elif key == "tech_category":
                val = classify_notice_tech(row).get("tech_category")
            else:
                val = row.get(key)

            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = border_all
            cell.alignment = left_align

            if key == "url" and val:
                cell.hyperlink = val
                cell.font = link_font
                cell.value = "View on World Bank"
                cell.alignment = center_align
            elif key in ("description", "overview", "requirements") and val:
                desc_text = str(val)
                if len(desc_text) > 500:
                    desc_text = desc_text[:500] + "..."
                cell.value = desc_text
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif key in ("notice_date", "submission_date", "awarded_date"):
                cell.alignment = center_align
                if val:
                    try:
                        if isinstance(val, str) and len(val) >= 10:
                            cell.value = datetime.strptime(val[:10], "%Y-%m-%d").date()
                    except Exception:
                        pass
            elif key in (
                "currency",
                "notice_type",
                "status",
                "won",
                "bidder_status",
                "is_tech",
                "opening_currency",
                "evaluated_bid_currency",
                "final_evaluation_currency",
                "winner_contract_currency",
            ):
                cell.alignment = center_align
                if key == "status" and val:
                    if val == "Active":
                        cell.fill = PatternFill("solid", fgColor="E8F5E8")
                    elif val == "Awarded":
                        cell.fill = PatternFill("solid", fgColor="E3F2FD")
                    elif val == "Closed":
                        cell.fill = PatternFill("solid", fgColor="FFF3E0")
                    elif val == "Cancelled":
                        cell.fill = PatternFill("solid", fgColor="FFEBEE")
            elif (
                key
                in (
                    "contract_amount",
                    "bid_price_at_opening",
                    "evaluated_bid_price",
                    "final_evaluation_price",
                    "winner_contract_amount",
                )
                and val is not None
            ):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                try:
                    amount = float(val)
                    if amount >= 1000000:
                        cell.number_format = '$#,##0.0,, "M"'
                    elif amount >= 1000:
                        cell.number_format = '$#,##0.0, "K"'
                    else:
                        cell.number_format = "$#,##0.00"
                except Exception:
                    cell.number_format = "#,##0.00"
            elif key == "contact_email" and val:
                cell.hyperlink = f"mailto:{val}"
                cell.font = Font(color="3498DB", underline="single", size=11, name="Calibri")

    ws2 = wb.create_sheet("Summary & Filters")
    ws2.column_dimensions["A"].width = 25
    ws2.column_dimensions["B"].width = 40
    ws2.column_dimensions["C"].width = 20

    title_cell = ws2.cell(row=1, column=1, value="World Bank Procurement Export Summary")
    title_cell.font = Font(bold=True, size=16, color="2C3E50", name="Calibri")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws2.merge_cells("A1:C1")

    export_type_name = (
        "Huzalink Export"
        if is_huzalink_export
        else ("Custom Bidder Export" if is_custom_bidder_export else "Normal Export")
    )
    ws2.cell(row=3, column=1, value=f"{export_type_name} Information").font = Font(
        bold=True, size=14, color="2C3E50", name="Calibri"
    )
    ws2.cell(row=3, column=1).fill = PatternFill("solid", fgColor="ECF0F1")
    ws2.merge_cells("A3:C3")

    info_data = [
        ("Export Type:", export_type_name),
        ("Export Date & Time:", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")),
        ("Total Rows:", len(rows)),
        ("Country Filter:", country or "All Countries"),
        ("Notice Type Filter:", notice_type or "All Types"),
        ("Status Filter:", status or "All Statuses"),
        ("Date Range From:", str(from_date) if from_date else "Not specified"),
        ("Date Range To:", str(to_date) if to_date else "Not specified"),
        ("Search Term:", search or "None"),
    ]
    if is_huzalink_export or is_custom_bidder_export:
        info_data.extend(
            [
                ("Selected Fields:", ", ".join(selected_fields)),
                ("Print Ready:", "Optimized for sharing"),
            ]
        )

    for r, (label, value) in enumerate(info_data, 4):
        label_cell = ws2.cell(row=r, column=1, value=label)
        label_cell.font = Font(bold=True, size=11, color="34495E", name="Calibri")
        label_cell.alignment = Alignment(horizontal="right", vertical="center")
        value_cell = ws2.cell(row=r, column=2, value=str(value))
        value_cell.font = Font(size=11, color="2C3E50", name="Calibri")
        value_cell.alignment = Alignment(horizontal="left", vertical="center")

    ws2.cell(row=len(info_data) + 6, column=1, value="Quick Statistics").font = Font(
        bold=True, size=14, color="2C3E50", name="Calibri"
    )
    ws2.cell(row=len(info_data) + 6, column=1).fill = PatternFill("solid", fgColor="ECF0F1")
    ws2.merge_cells(f"A{len(info_data) + 6}:C{len(info_data) + 6}")

    active_count = sum(1 for r in rows if r.get("status") == "Active")
    awarded_count = sum(1 for r in rows if r.get("status") == "Awarded")
    stats_data = [
        ("Active Opportunities:", active_count),
        ("Awarded Opportunities:", awarded_count),
        ("Countries Represented:", len(set(r.get("country") for r in rows if r.get("country")))),
        ("Average Contract Amount:", "N/A"),
    ]
    for r, (label, value) in enumerate(stats_data, len(info_data) + 7):
        label_cell = ws2.cell(row=r, column=1, value=label)
        label_cell.font = Font(bold=True, size=11, color="34495E", name="Calibri")
        label_cell.alignment = Alignment(horizontal="right", vertical="center")
        value_cell = ws2.cell(row=r, column=2, value=str(value))
        value_cell.font = Font(size=11, color="2C3E50", name="Calibri")
        value_cell.alignment = Alignment(horizontal="left", vertical="center")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M")
    if is_huzalink_export:
        parts = ["Huzalink_Procurement", timestamp]
    elif is_custom_bidder_export:
        parts = ["WB_Procurement_Bidders", timestamp]
    else:
        parts = ["WB_Procurement", timestamp]
    if country:
        parts.append(country.replace(" ", "_")[:20])
    if notice_type:
        parts.append(notice_type[:15])
    if status:
        parts.append(status[:15])
    filename = "_".join(parts) + ".xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/csv")
def export_csv(
    country: str | None = None,
    notice_type: str | None = None,
    status: str | None = None,
    search: str | None = None,
    from_date: str | None = None,
    to_date: str | None = None,
    fields: str | None = None,
):
    where, params = build_where(
        country=country, notice_type=notice_type, status=status, search=search, from_date=from_date, to_date=to_date
    )
    is_custom_bidder_export = bool(fields)
    if is_custom_bidder_export:
        selected_fields = resolve_custom_bidder_export_fields(fields)
        rows = fetch_custom_bidder_export_rows(where, params, selected_fields)
    else:
        selected_fields = resolve_export_fields(fields)
        rows = fetch_export_rows(where, params, selected_fields)

    output = io.StringIO()
    writer = csv.writer(output)
    field_config = CUSTOM_BIDDER_EXPORT_FIELD_CONFIG if is_custom_bidder_export else EXPORT_FIELD_CONFIG
    headers = [field_config[field][0] for field in selected_fields if field in field_config]
    writer.writerow(headers)
    for row in rows:
        values = []
        for field in selected_fields:
            if not is_custom_bidder_export and field == "bidders":
                values.append(
                    extract_bidders_from_description(row.get("description")) if row.get("description") else None
                )
            elif field == "overview":
                values.append(extract_notice_overview(row.get("description")))
            elif field == "requirements":
                values.append(extract_notice_requirements(row.get("description")))
            elif field == "is_tech":
                values.append("Yes" if classify_notice_tech(row).get("is_tech") else "No")
            elif field == "tech_category":
                values.append(classify_notice_tech(row).get("tech_category"))
            else:
                values.append(row.get(field_config[field][1]))
        writer.writerow(values)

    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M")
    filename = f"{'WB_Procurement_Bidders' if is_custom_bidder_export else 'WB_Procurement'}_{timestamp}.csv"
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
