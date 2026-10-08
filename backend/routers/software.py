import io
import json
from datetime import datetime
from typing import Any

import openpyxl
import requests
from auth import require_operator
from db import db, q
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from ratelimit import STRICT_LIMITER, rate_limited
from services.software_intelligence import classify_software_opportunity_with_gemini
from services.tech import build_tech_notice_condition

router = APIRouter(prefix="/api/software-opportunities", tags=["software"])


class SoftwareClassificationReview(BaseModel):
    is_software_related: bool
    classification: str = "software"
    classification_status: str = "classified"
    technology_related: bool = True
    primary_deliverable: str | None = None
    work_type: str = "other"
    product_category: str | None = None
    product_name: str | None = None
    confidence: str = "medium"
    reason: str | None = None


def ensure_software_intelligence_schema():
    return None


def save_software_classification(notice_id: str, result: dict, reviewed: bool = False):
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT software_reviewed FROM procurement_notices WHERE id = %s",
                [notice_id],
            )
            existing = cur.fetchone()
            if not existing:
                raise HTTPException(status_code=404, detail="Notice not found")
            if existing["software_reviewed"] and not reviewed:
                return
            cur.execute(
                """
                UPDATE procurement_notices
                SET is_software_related = %s,
                    classification = %s,
                    classification_status = %s,
                    technology_related = %s,
                    primary_deliverable = %s,
                    software_work_type = %s,
                    software_product_category = %s,
                    software_subcategory = %s,
                    software_product_name = %s,
                    software_confidence = %s,
                    software_reason = %s,
                    classification_evidence = %s::jsonb,
                    classifier_name = %s,
                    classifier_version = %s,
                    prompt_version = %s,
                    classification_model = %s,
                    software_classification_source = %s,
                    software_classified_at = NOW(),
                    software_reviewed = %s
                WHERE id = %s
                RETURNING id
            """,
                [
                    result["is_software_related"],
                    result["classification"],
                    result["classification_status"],
                    result["technology_related"],
                    result["primary_deliverable"],
                    result["work_type"],
                    result["product_category"],
                    result["product_subcategory"],
                    result["product_name"],
                    result["confidence"],
                    result["reason"],
                    json.dumps(result.get("evidence", [])),
                    result["classifier_name"],
                    result["classifier_version"],
                    result["prompt_version"],
                    result["classification_model"],
                    result.get("source", "gemini"),
                    reviewed,
                    notice_id,
                ],
            )
            row = cur.fetchone()
            if row:
                cur.execute(
                    """
                    INSERT INTO notice_classifications (
                        notice_id, classification, classification_status, is_software_related,
                        technology_related, primary_deliverable, software_work_type,
                        software_product_category, software_subcategory, software_product_name,
                        confidence, reason, evidence, rule_signals, classifier_name,
                        classifier_version, prompt_version, classification_model, source, reviewed
                    )
                    SELECT id, classification, classification_status, is_software_related,
                           technology_related, primary_deliverable, software_work_type,
                           software_product_category, software_subcategory, software_product_name,
                           software_confidence, software_reason,
                           COALESCE(classification_evidence, '[]'::jsonb), %s::jsonb,
                           classifier_name, classifier_version, prompt_version,
                           classification_model, software_classification_source, %s
                    FROM procurement_notices
                    WHERE id = %s
                    """,
                    [json.dumps(result.get("rule_signals", {})), reviewed, notice_id],
                )
        conn.commit()
    if not row:
        raise HTTPException(status_code=404, detail="Notice not found")


# ── Endpoints ─────────────────────────────────────────────────────────────────


@router.get("")
def get_software_opportunities(
    country: str | None = Query(None),
    product_category: str | None = Query(None),
    classification: str | None = Query(None),
    classification_status: str | None = Query(None),
    technology_related: bool | None = Query(None),
    min_confidence: str | None = Query(None, regex="^(high|medium|low)$"),
    unreviewed_only: bool = Query(False),
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100),
):
    ensure_software_intelligence_schema()
    filters = ["(pn.classification IS NOT NULL OR pn.is_software_related IS TRUE)"]
    params: list[Any] = []
    if country:
        filters.append("pn.country = %s")
        params.append(country)
    if product_category:
        filters.append("pn.software_product_category ILIKE %s")
        params.append(f"%{product_category}%")
    if classification:
        filters.append("pn.classification = %s")
        params.append(classification)
    if classification_status:
        filters.append("pn.classification_status = %s")
        params.append(classification_status)
    if technology_related is not None:
        filters.append("pn.technology_related IS %s")
        params.append(technology_related)
    if min_confidence:
        levels = {"high": ["high"], "medium": ["high", "medium"], "low": ["high", "medium", "low"]}
        filters.append("pn.software_confidence = ANY(%s)")
        params.append(levels[min_confidence])
    if unreviewed_only:
        filters.append("pn.software_reviewed IS FALSE")
    where = " AND ".join(filters)
    offset = (page - 1) * page_size
    rows = q(
        f"""
        SELECT pn.id, pn.title, pn.country, pn.notice_type, pn.status, pn.notice_date::text,
               pn.submission_date::text, pn.url, pn.software_work_type,
               pn.software_product_category, pn.software_product_name,
               pn.software_subcategory, pn.software_confidence, pn.software_reason,
               pn.classification, pn.classification_status, pn.technology_related,
               pn.primary_deliverable, pn.classification_evidence, pn.software_reviewed,
               COUNT(DISTINCT COALESCE(NULLIF(related.borrower_bid_reference, ''), related.id))
                   FILTER (WHERE related.is_software_related IS TRUE) AS category_demand_count,
               STRING_AGG(DISTINCT related.country, ', ' ORDER BY related.country)
                   FILTER (WHERE related.is_software_related IS TRUE AND related.country IS NOT NULL) AS demand_countries
        FROM procurement_notices pn
        LEFT JOIN procurement_notices related
          ON related.is_software_related IS TRUE
         AND related.software_product_category = pn.software_product_category
        WHERE {where}
        GROUP BY pn.id
        ORDER BY category_demand_count DESC, pn.notice_date DESC NULLS LAST
        LIMIT %s OFFSET %s
    """,
        params + [page_size, offset],
    )
    total = q(f"SELECT COUNT(*) AS cnt FROM procurement_notices pn WHERE {where}", params)[0]["cnt"]
    categories = q("""
        SELECT software_product_category AS category, COUNT(*) AS demand_count
        FROM procurement_notices
        WHERE is_software_related IS TRUE AND software_product_category IS NOT NULL
        GROUP BY software_product_category
        ORDER BY demand_count DESC, software_product_category
        LIMIT 100
    """)
    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "data": [dict(row) for row in rows],
        "categories": [dict(row) for row in categories],
    }


@router.post("/{notice_id}/classify", dependencies=[Depends(rate_limited(STRICT_LIMITER))])
def classify_software_opportunity(notice_id: str, _: dict = Depends(require_operator)):
    ensure_software_intelligence_schema()
    rows = q("SELECT * FROM procurement_notices WHERE id = %s", [notice_id])
    if not rows:
        raise HTTPException(status_code=404, detail="Notice not found")
    try:
        result = classify_software_opportunity_with_gemini(dict(rows[0]))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"Classification provider failed: {exc}") from exc
    save_software_classification(notice_id, result)
    return {"notice_id": notice_id, **result, "source": "gemini"}


@router.post("/classify-pending", dependencies=[Depends(rate_limited(STRICT_LIMITER))])
def classify_pending_software_opportunities(limit: int = Query(10, ge=1, le=25), _: dict = Depends(require_operator)):
    ensure_software_intelligence_schema()
    tech_condition, tech_params = build_tech_notice_condition("pn")
    notices = q(
        f"""
        SELECT * FROM procurement_notices pn
        WHERE pn.is_software_related IS NULL
        ORDER BY CASE WHEN {tech_condition} THEN 0 ELSE 1 END,
                 pn.notice_date DESC NULLS LAST, pn.fetched_at DESC NULLS LAST
        LIMIT %s
    """,
        tech_params + [limit],
    )
    completed, failed = [], []
    for notice in notices:
        try:
            result = classify_software_opportunity_with_gemini(dict(notice))
            save_software_classification(notice["id"], result)
            completed.append({"notice_id": notice["id"], "is_software_related": result["is_software_related"]})
        except Exception as exc:
            failed.append({"notice_id": notice["id"], "error": str(exc)})
    return {"requested": len(notices), "classified": completed, "failed": failed}


@router.put("/{notice_id}/review")
def review_software_opportunity(
    notice_id: str, body: SoftwareClassificationReview, _: dict = Depends(require_operator)
):
    ensure_software_intelligence_schema()
    result = body.model_dump()
    if result["work_type"] not in {"new_build", "enhancement", "integration", "implementation", "support", "other"}:
        raise HTTPException(status_code=422, detail="Invalid work_type")
    if result["confidence"] not in {"high", "medium", "low"}:
        raise HTTPException(status_code=422, detail="Invalid confidence")
    if result["classification"] not in {
        "software",
        "hardware",
        "civil_works",
        "consulting",
        "training",
        "infrastructure",
        "telecommunications",
        "data_services",
        "general_services",
        "goods_supply",
        "mixed",
        "unclear",
    }:
        raise HTTPException(status_code=422, detail="Invalid classification")
    if not result["is_software_related"]:
        result.update({"work_type": "other", "product_category": None, "product_name": None})
    existing = q("SELECT * FROM procurement_notices WHERE id = %s", [notice_id])
    if not existing:
        raise HTTPException(status_code=404, detail="Notice not found")
    current = dict(existing[0])
    result.update(
        {
            "product_subcategory": current.get("software_subcategory"),
            "evidence": current.get("classification_evidence") or [],
            "rule_signals": {},
            "classifier_name": "human_review",
            "classifier_version": "1",
            "prompt_version": None,
            "classification_model": None,
            "source": "human",
        }
    )
    save_software_classification(notice_id, result, reviewed=True)
    return {"status": "ok", "notice_id": notice_id, **result, "source": "human"}


@router.get("/export")
def export_software_opportunities(
    country: str | None = Query(None),
    product_category: str | None = Query(None),
):
    ensure_software_intelligence_schema()
    filters = ["pn.is_software_related IS TRUE"]
    params: list[Any] = []
    if country:
        filters.append("pn.country = %s")
        params.append(country)
    if product_category:
        filters.append("pn.software_product_category ILIKE %s")
        params.append(f"%{product_category}%")
    where = " AND ".join(filters)
    rows = q(
        f"""
        SELECT pn.id, pn.title, pn.country, pn.notice_type, pn.status, pn.notice_date::text,
               pn.submission_date::text, pn.url, pn.software_work_type,
               pn.software_product_category, pn.software_product_name,
               pn.software_confidence, pn.software_reason,
               COUNT(DISTINCT COALESCE(NULLIF(related.borrower_bid_reference, ''), related.id))
                   FILTER (WHERE related.is_software_related IS TRUE) AS demand_count,
               STRING_AGG(DISTINCT related.country, ', ' ORDER BY related.country)
                   FILTER (WHERE related.is_software_related IS TRUE AND related.country IS NOT NULL) AS demand_countries,
               MIN(related.notice_date)::text FILTER (WHERE related.is_software_related IS TRUE) AS first_seen,
               MAX(related.notice_date)::text FILTER (WHERE related.is_software_related IS TRUE) AS last_seen
        FROM procurement_notices pn
        LEFT JOIN procurement_notices related
          ON related.is_software_related IS TRUE
         AND related.software_product_category = pn.software_product_category
        WHERE {where}
        GROUP BY pn.id
        ORDER BY demand_count DESC, pn.notice_date DESC NULLS LAST
    """,
        params,
    )
    history_where = where.replace("pn.", "base_notice.")
    history = q(
        f"""
        SELECT base_notice.id AS opportunity_id, base_notice.title AS opportunity_title,
               base_notice.software_product_category, historical.id AS historical_notice_id,
               historical.title AS historical_title, historical.country,
               historical.notice_type, historical.status, historical.notice_date::text,
               historical.url, historical.borrower_bid_reference
        FROM procurement_notices base_notice
        JOIN procurement_notices historical
          ON historical.is_software_related IS TRUE
         AND historical.software_product_category = base_notice.software_product_category
        WHERE {history_where}
        ORDER BY historical.software_product_category, historical.notice_date DESC NULLS LAST
    """,
        params,
    )

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Ranked Opportunities"
    columns = [
        ("Opportunity", "title"),
        ("Product category", "software_product_category"),
        ("Product name", "software_product_name"),
        ("Work type", "software_work_type"),
        ("Demand count", "demand_count"),
        ("Countries requested in", "demand_countries"),
        ("First seen", "first_seen"),
        ("Last seen", "last_seen"),
        ("Country", "country"),
        ("Notice date", "notice_date"),
        ("Confidence", "software_confidence"),
        ("Why classified", "software_reason"),
        ("Source URL", "url"),
    ]
    sheet.append([label for label, _ in columns])
    for row in rows:
        sheet.append([row.get(key) for _, key in columns])
    history_sheet = workbook.create_sheet("Demand History")
    history_columns = [
        ("Opportunity", "opportunity_title"),
        ("Product category", "software_product_category"),
        ("Historical tender", "historical_title"),
        ("Country", "country"),
        ("Notice type", "notice_type"),
        ("Status", "status"),
        ("Notice date", "notice_date"),
        ("Reference", "borrower_bid_reference"),
        ("Source URL", "url"),
    ]
    history_sheet.append([label for label, _ in history_columns])
    for row in history:
        history_sheet.append([row.get(key) for _, key in history_columns])
    for ws in (sheet, history_sheet):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for column in ws.columns:
            letter = get_column_letter(column[0].column)
            ws.column_dimensions[letter].width = min(
                max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 55
            )
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="Software_Opportunity_Demand_{timestamp}.xlsx"'},
    )
